"""Knowledge Base indexing tasks."""
from __future__ import annotations
import asyncio
import hashlib
import logging
import mimetypes
from pathlib import Path
import httpx
from celery import shared_task
from sqlalchemy import select, delete
from sqlalchemy.ext.asyncio import create_async_engine, async_sessionmaker
from app.db.models import DocumentStatus

logger = logging.getLogger(__name__)

CHUNK_SIZE = 512      # tokens approx (chars / 4)
CHUNK_OVERLAP = 64
EMBED_MODEL = "nomic-embed-text"
EMBED_DIM = 768


def _get_session() -> async_sessionmaker:
    from app.config import get_settings
    s = get_settings()
    engine = create_async_engine(s.database_url, pool_pre_ping=True)
    return async_sessionmaker(engine, expire_on_commit=False)


def _extract_text(path: str, mime: str) -> str:
    """Extract plain text from file. Returns empty string on failure."""
    try:
        if mime == "application/pdf" or path.endswith(".pdf"):
            import pdfplumber
            with pdfplumber.open(path) as pdf:
                return "\n".join(p.extract_text() or "" for p in pdf.pages)
        elif mime in ("application/vnd.openxmlformats-officedocument.wordprocessingml.document",) or path.endswith(".docx"):
            from docx import Document as DocxDoc
            doc = DocxDoc(path)
            return "\n".join(p.text for p in doc.paragraphs)
        elif mime in ("application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",) or path.endswith(".xlsx"):
            import openpyxl
            wb = openpyxl.load_workbook(path, data_only=True)
            lines = []
            for ws in wb.worksheets:
                for row in ws.iter_rows(values_only=True):
                    lines.append(" | ".join(str(c) for c in row if c is not None))
            return "\n".join(lines)
        else:
            # plain text fallback
            with open(path, "r", encoding="utf-8", errors="replace") as f:
                return f.read()
    except Exception as e:
        logger.error(f"Text extraction failed for {path}: {e}")
        return ""


def _chunk_text(text: str, size: int = CHUNK_SIZE * 4, overlap: int = CHUNK_OVERLAP * 4) -> list[str]:
    """Split text into overlapping character chunks."""
    if not text.strip():
        return []
    chunks = []
    start = 0
    while start < len(text):
        end = start + size
        chunks.append(text[start:end].strip())
        start += size - overlap
    return [c for c in chunks if c]


async def _embed_text(texts: list[str], ollama_url: str) -> list[list[float]]:
    """Get embeddings from Ollama. Returns list of float vectors."""
    async with httpx.AsyncClient(timeout=60.0) as client:
        embeddings = []
        for text in texts:
            r = await client.post(
                f"{ollama_url}/api/embeddings",
                json={"model": EMBED_MODEL, "prompt": text},
            )
            r.raise_for_status()
            embeddings.append(r.json()["embedding"])
        return embeddings


async def _index_document_async(doc_id: str, org_id: str, file_path: str, mime: str) -> dict:
    """Full indexing pipeline for one document."""
    from app.db.models import Document, DocumentChunk
    from app.config import get_settings
    import uuid

    settings = get_settings()
    SessionLocal = _get_session()

    async with SessionLocal() as db:
        # Mark as processing
        doc = await db.get(Document, uuid.UUID(doc_id))
        if not doc:
            return {"error": "document not found"}
        doc.status = DocumentStatus.processing.value
        await db.commit()

        try:
            # 1. Extract text
            text = _extract_text(file_path, mime)
            if not text.strip():
                doc.status = DocumentStatus.failed.value
                doc.metadata_json = {"error": "no text extracted"}
                await db.commit()
                return {"error": "no text extracted"}

            # 2. Chunk
            chunks = _chunk_text(text)
            if not chunks:
                doc.status = DocumentStatus.failed.value
                await db.commit()
                return {"error": "chunking produced no chunks"}

            # 3. Delete existing chunks for this doc
            await db.execute(delete(DocumentChunk).where(DocumentChunk.document_id == doc.id))

            # 4. Embed
            embeddings = await _embed_text(chunks, settings.ollama_base_url)

            # 5. Store chunks
            for i, (chunk_text, embedding) in enumerate(zip(chunks, embeddings)):
                chunk = DocumentChunk(
                    document_id=doc.id,
                    org_id=uuid.UUID(org_id),
                    chunk_index=i,
                    content=chunk_text,
                    embedding=embedding,
                )
                db.add(chunk)

            # 6. Mark indexed
            doc.status = DocumentStatus.indexed.value
            doc.total_chunks = len(chunks)
            await db.commit()
            return {"chunks": len(chunks), "status": DocumentStatus.indexed.value}

        except Exception as e:
            logger.error(f"Indexing failed for doc {doc_id}: {e}")
            async with SessionLocal() as db2:
                doc2 = await db2.get(Document, uuid.UUID(doc_id))
                if doc2:
                    doc2.status = DocumentStatus.failed.value
                    doc2.metadata_json = {"error": str(e)}
                    await db2.commit()
            return {"error": str(e)}


@shared_task(bind=True, name="app.tasks.kb.index_document", max_retries=3, default_retry_delay=30)
def index_document(self, doc_id: str, org_id: str, file_path: str, mime: str = "text/plain"):
    """Celery task: index a single document into the knowledge base."""
    try:
        loop = asyncio.new_event_loop()
        asyncio.set_event_loop(loop)
        result = loop.run_until_complete(_index_document_async(doc_id, org_id, file_path, mime))
        loop.close()
        return result
    except Exception as exc:
        raise self.retry(exc=exc)


@shared_task(name="app.tasks.kb.scan_watch_paths")
def scan_watch_paths():
    """Celery beat task: scan all enabled KbWatchPath entries."""
    from app.db.models import KbWatchPath, Document

    async def _scan():
        from app.config import get_settings
        get_settings()
        SessionLocal = _get_session()
        async with SessionLocal() as db:
            result = await db.execute(
                select(KbWatchPath).where(KbWatchPath.is_enabled)
            )
            paths = result.scalars().all()
            queued = 0
            for wp in paths:
                if wp.path_type == "local":
                    p = Path(wp.path)
                    if not p.exists():
                        continue
                    for fpath in p.rglob("*"):
                        if fpath.suffix.lower() in (".pdf", ".docx", ".xlsx", ".xls", ".dxf", ".eml", ".msg", ".txt", ".md", ".csv"):
                            content_hash = hashlib.sha256(fpath.read_bytes()).hexdigest()
                            # Check if already indexed with same hash
                            existing = await db.execute(
                                select(Document).where(
                                    Document.org_id == wp.org_id,
                                    Document.source_path == str(fpath),
                                    Document.content_hash == content_hash,
                                    Document.status == DocumentStatus.indexed.value,
                                )
                            )
                            if existing.scalar_one_or_none():
                                continue
                            inferred_mime = mimetypes.guess_type(str(fpath))[0] or "application/octet-stream"
                            doc = Document(
                                org_id=wp.org_id,
                                title=fpath.name,
                                filename=fpath.name,
                                source_path=str(fpath),
                                content_hash=content_hash,
                                status=DocumentStatus.pending.value,
                                mime_type=inferred_mime,
                                meta_json={},
                            )
                            db.add(doc)
                            await db.flush()
                            index_document.delay(str(doc.id), str(wp.org_id), str(fpath))
                            queued += 1
            await db.commit()
            return {"queued": queued}

    loop = asyncio.new_event_loop()
    asyncio.set_event_loop(loop)
    result = loop.run_until_complete(_scan())
    loop.close()
    return result
