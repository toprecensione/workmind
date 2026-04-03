"""
WorkMind Document Parser
CONFIDENTIAL - PRIVATE REPOSITORY - NOT FOR PUBLIC DISTRIBUTION

Estrae testo strutturato da: PDF, DOCX, XLSX/XLS, CSV, TXT, EML, MSG, DXF/DWG.
Ogni formato viene tentato con la libreria dedicata; fallback su estrazione raw.
"""

from __future__ import annotations

import csv
import hashlib
import io
import os
import re
import zipfile
from dataclasses import dataclass, field
from datetime import datetime, timezone
from pathlib import Path
from typing import Optional

from logging_system import get_logger, LogStatus, LogAction
from storage.redis_store import get_store

log = get_logger("connectors.document_parser")

_SUPPORTED_EXTENSIONS = {
    ".pdf", ".docx", ".doc", ".xlsx", ".xls",
    ".csv", ".txt", ".eml", ".msg", ".dxf", ".dwg",
}


@dataclass
class ParsedDocument:
    path: str
    filename: str
    extension: str
    text: str                             # testo estratto (plain)
    pages: int = 0
    tables: list[list[list[str]]] = field(default_factory=list)  # tabelle estratte
    metadata: dict = field(default_factory=dict)
    sha256: str = ""
    parsed_at: str = ""
    error: Optional[str] = None

    @property
    def is_ok(self) -> bool:
        return self.error is None and bool(self.text.strip())

    @property
    def word_count(self) -> int:
        return len(self.text.split())

    @property
    def preview(self) -> str:
        return self.text[:500].replace("\n", " ").strip()


class DocumentParser:
    """
    Parser universale di documenti aziendali.
    Usa librerie opzionali — se non installate, avverte e salta il formato.
    """

    def __init__(self) -> None:
        self._store = get_store()

    def parse(self, path: str | Path) -> ParsedDocument:
        """
        Parsa un documento. Se già in cache (stesso hash), ritorna il risultato
        memorizzato senza ri-parsare.
        """
        p = Path(path)
        ext = p.suffix.lower()

        if ext not in _SUPPORTED_EXTENSIONS:
            return self._error_doc(p, f"Estensione non supportata: {ext}")

        if not p.exists():
            return self._error_doc(p, "File non trovato")

        sha = self._sha256(p)
        cached = self._store.get_cached_document(sha)
        if cached:
            log.debug(f"Cache hit: {p.name}", action=LogAction.SCAN)
            return ParsedDocument(**cached)

        doc = self._dispatch(p, ext, sha)
        if doc.is_ok:
            self._store.cache_document(sha, {
                "path": doc.path, "filename": doc.filename, "extension": doc.extension,
                "text": doc.text[:10000],  # salva max 10k caratteri in cache
                "pages": doc.pages, "tables": [], "metadata": doc.metadata,
                "sha256": doc.sha256, "parsed_at": doc.parsed_at, "error": None,
            })

        return doc

    def parse_batch(self, paths: list[str | Path]) -> list[ParsedDocument]:
        results = []
        for p in paths:
            try:
                results.append(self.parse(p))
            except Exception as exc:
                results.append(self._error_doc(Path(p), str(exc)))
        return results

    # ── Dispatch per formato ──────────────────────────────────────────────────

    def _dispatch(self, p: Path, ext: str, sha: str) -> ParsedDocument:
        handlers = {
            ".pdf":  self._parse_pdf,
            ".docx": self._parse_docx,
            ".doc":  self._parse_doc,
            ".xlsx": self._parse_excel,
            ".xls":  self._parse_excel,
            ".csv":  self._parse_csv,
            ".txt":  self._parse_txt,
            ".eml":  self._parse_eml,
            ".msg":  self._parse_msg,
            ".dxf":  self._parse_dxf,
            ".dwg":  self._parse_dwg,
        }
        handler = handlers.get(ext, self._parse_txt)
        doc = handler(p)
        doc.sha256 = sha
        doc.parsed_at = datetime.now(timezone.utc).isoformat()
        return doc

    # ── PDF ───────────────────────────────────────────────────────────────────

    def _parse_pdf(self, p: Path) -> ParsedDocument:
        try:
            import pdfplumber
            text_parts = []
            tables = []
            n_pages = 0
            with pdfplumber.open(str(p)) as pdf:
                n_pages = len(pdf.pages)
                for page in pdf.pages:
                    text_parts.append(page.extract_text() or "")
                    for tbl in page.extract_tables() or []:
                        if tbl:
                            tables.append([[str(c) if c else "" for c in row] for row in tbl])
            return self._make_doc(p, "\n".join(text_parts), pages=n_pages, tables=tables)
        except ImportError:
            pass

        # Fallback: pypdf2
        try:
            import PyPDF2
            text_parts = []
            with open(p, "rb") as f:
                reader = PyPDF2.PdfReader(f)
                for page in reader.pages:
                    text_parts.append(page.extract_text() or "")
            return self._make_doc(p, "\n".join(text_parts), pages=len(reader.pages))
        except ImportError:
            return self._error_doc(p, "Installa pdfplumber: pip install pdfplumber")

    # ── DOCX ──────────────────────────────────────────────────────────────────

    def _parse_docx(self, p: Path) -> ParsedDocument:
        try:
            import docx
            doc = docx.Document(str(p))
            paragraphs = [para.text for para in doc.paragraphs]
            tables = []
            for tbl in doc.tables:
                table_data = []
                for row in tbl.rows:
                    table_data.append([cell.text for cell in row.cells])
                tables.append(table_data)
            text = "\n".join(paragraphs)
            meta = {
                "author": doc.core_properties.author or "",
                "created": str(doc.core_properties.created or ""),
                "modified": str(doc.core_properties.modified or ""),
            }
            return self._make_doc(p, text, tables=tables, metadata=meta)
        except ImportError:
            # Fallback: estrazione raw da XML dentro il zip
            try:
                with zipfile.ZipFile(p, "r") as z:
                    with z.open("word/document.xml") as f:
                        content = f.read().decode("utf-8")
                text = re.sub(r"<[^>]+>", " ", content)
                text = re.sub(r"\s+", " ", text).strip()
                return self._make_doc(p, text)
            except Exception as exc:
                return self._error_doc(p, f"Errore DOCX: {exc}")

    def _parse_doc(self, p: Path) -> ParsedDocument:
        # Legacy .doc: richiede antiword o catdoc su Ubuntu
        try:
            import subprocess
            result = subprocess.run(
                ["antiword", str(p)], capture_output=True, text=True, timeout=30
            )
            if result.returncode == 0:
                return self._make_doc(p, result.stdout)
        except (FileNotFoundError, Exception):
            pass
        return self._error_doc(p, "File .doc legacy: installa antiword su Ubuntu")

    # ── Excel ─────────────────────────────────────────────────────────────────

    def _parse_excel(self, p: Path) -> ParsedDocument:
        try:
            import openpyxl
            wb = openpyxl.load_workbook(str(p), read_only=True, data_only=True)
            text_parts = []
            tables = []
            for sheet in wb.worksheets:
                text_parts.append(f"[Foglio: {sheet.title}]")
                sheet_rows = []
                for row in sheet.iter_rows(values_only=True):
                    cells = [str(c) if c is not None else "" for c in row]
                    if any(c.strip() for c in cells):
                        text_parts.append("\t".join(cells))
                        sheet_rows.append(cells)
                if sheet_rows:
                    tables.append(sheet_rows)
            wb.close()
            return self._make_doc(p, "\n".join(text_parts), tables=tables)
        except ImportError:
            pass

        # Fallback: openpyxl non disponibile, tenta xlrd per .xls
        try:
            import xlrd
            wb = xlrd.open_workbook(str(p))
            text_parts = []
            for sheet in wb.sheets():
                text_parts.append(f"[Foglio: {sheet.name}]")
                for row_idx in range(sheet.nrows):
                    cells = [str(sheet.cell_value(row_idx, col)) for col in range(sheet.ncols)]
                    text_parts.append("\t".join(cells))
            return self._make_doc(p, "\n".join(text_parts))
        except ImportError:
            return self._error_doc(p, "Installa openpyxl: pip install openpyxl")

    # ── CSV ───────────────────────────────────────────────────────────────────

    def _parse_csv(self, p: Path) -> ParsedDocument:
        try:
            encoding = self._detect_encoding(p)
            rows = []
            with open(p, newline="", encoding=encoding, errors="replace") as f:
                # Tenta di rilevare il delimitatore
                sample = f.read(4096)
                f.seek(0)
                dialect = csv.Sniffer().sniff(sample, delimiters=";,\t|")
                reader = csv.reader(f, dialect)
                for row in reader:
                    rows.append(row)

            text = "\n".join([";".join(row) for row in rows[:200]])  # max 200 righe in testo
            return self._make_doc(p, text, tables=[rows] if rows else [])
        except Exception as exc:
            # Fallback: leggi come testo puro
            try:
                text = p.read_text(encoding="utf-8", errors="replace")
                return self._make_doc(p, text)
            except Exception:
                return self._error_doc(p, f"Errore CSV: {exc}")

    # ── TXT ───────────────────────────────────────────────────────────────────

    def _parse_txt(self, p: Path) -> ParsedDocument:
        try:
            encoding = self._detect_encoding(p)
            text = p.read_text(encoding=encoding, errors="replace")
            return self._make_doc(p, text)
        except Exception as exc:
            return self._error_doc(p, f"Errore TXT: {exc}")

    # ── EML (email) ───────────────────────────────────────────────────────────

    def _parse_eml(self, p: Path) -> ParsedDocument:
        import email as email_lib
        try:
            msg = email_lib.message_from_bytes(p.read_bytes())
            subject  = msg.get("Subject", "")
            sender   = msg.get("From", "")
            date_str = msg.get("Date", "")
            body_parts = []
            for part in msg.walk():
                ctype = part.get_content_type()
                if ctype in ("text/plain", "text/html"):
                    payload = part.get_payload(decode=True)
                    if payload:
                        text = payload.decode(part.get_content_charset() or "utf-8", errors="replace")
                        if ctype == "text/html":
                            text = re.sub(r"<[^>]+>", " ", text)
                        body_parts.append(text)
            body = "\n".join(body_parts)
            full_text = f"Oggetto: {subject}\nDa: {sender}\nData: {date_str}\n\n{body}"
            return self._make_doc(p, full_text, metadata={"subject": subject, "from": sender, "date": date_str})
        except Exception as exc:
            return self._error_doc(p, f"Errore EML: {exc}")

    # ── MSG (Outlook) ─────────────────────────────────────────────────────────

    def _parse_msg(self, p: Path) -> ParsedDocument:
        try:
            import extract_msg
            msg = extract_msg.Message(str(p))
            text = (
                f"Oggetto: {msg.subject or ''}\n"
                f"Da: {msg.sender or ''}\n"
                f"Data: {msg.date or ''}\n\n"
                f"{msg.body or ''}"
            )
            return self._make_doc(p, text, metadata={
                "subject": msg.subject or "",
                "from": msg.sender or "",
                "date": str(msg.date or ""),
            })
        except ImportError:
            return self._error_doc(p, "Installa extract-msg: pip install extract-msg")
        except Exception as exc:
            return self._error_doc(p, f"Errore MSG: {exc}")

    # ── DXF/DWG (AutoCAD) ────────────────────────────────────────────────────

    def _parse_dxf(self, p: Path) -> ParsedDocument:
        try:
            import ezdxf
            doc = ezdxf.readfile(str(p))
            texts = []
            layers = set()
            blocks = []
            msp = doc.modelspace()
            for entity in msp:
                layers.add(entity.dxf.layer if hasattr(entity.dxf, "layer") else "")
                if entity.dxftype() in ("TEXT", "MTEXT"):
                    t = getattr(entity.dxf, "text", None) or getattr(entity.dxf, "insert", None)
                    if t:
                        texts.append(str(t))
            for block in doc.blocks:
                blocks.append(block.name)
            text = "\n".join(texts)
            meta = {
                "layers": sorted(layers - {""}),
                "blocks": blocks,
                "dxfversion": doc.dxfversion,
            }
            return self._make_doc(p, text, metadata=meta)
        except ImportError:
            return self._error_doc(p, "Installa ezdxf: pip install ezdxf")
        except Exception as exc:
            return self._error_doc(p, f"Errore DXF: {exc}")

    def _parse_dwg(self, p: Path) -> ParsedDocument:
        # DWG binario: prova conversione a DXF con LibreDWG se disponibile
        try:
            import subprocess
            dxf_path = p.with_suffix(".dxf")
            result = subprocess.run(
                ["dwg2dxf", str(p), "-o", str(dxf_path)],
                capture_output=True, timeout=60
            )
            if result.returncode == 0 and dxf_path.exists():
                doc = self._parse_dxf(dxf_path)
                try:
                    dxf_path.unlink()
                except Exception:
                    pass
                return doc
        except (FileNotFoundError, Exception):
            pass
        return self._error_doc(p, "File DWG: installa LibreDWG (dwg2dxf) per l'estrazione")

    # ── Utility ───────────────────────────────────────────────────────────────

    def _make_doc(
        self,
        p: Path,
        text: str,
        pages: int = 1,
        tables: list | None = None,
        metadata: dict | None = None,
    ) -> ParsedDocument:
        return ParsedDocument(
            path=str(p),
            filename=p.name,
            extension=p.suffix.lower(),
            text=text.strip(),
            pages=pages,
            tables=tables or [],
            metadata=metadata or {},
        )

    def _error_doc(self, p: Path, error: str) -> ParsedDocument:
        log.warning(
            f"Parser: {error} — file: {p.name}",
            action=LogAction.SCAN, status=LogStatus.WARNING,
        )
        return ParsedDocument(
            path=str(p),
            filename=p.name,
            extension=p.suffix.lower(),
            text="",
            error=error,
            parsed_at=datetime.now(timezone.utc).isoformat(),
        )

    @staticmethod
    def _sha256(p: Path) -> str:
        h = hashlib.sha256()
        with open(p, "rb") as f:
            for chunk in iter(lambda: f.read(65536), b""):
                h.update(chunk)
        return h.hexdigest()

    @staticmethod
    def _detect_encoding(p: Path) -> str:
        try:
            import chardet
            raw = p.read_bytes()[:10000]
            result = chardet.detect(raw)
            return result.get("encoding") or "utf-8"
        except ImportError:
            return "utf-8"
